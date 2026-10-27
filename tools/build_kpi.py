# Membuat assets/js/kpi-data.js dari KPI_OBO_*.xlsx (sheet "KPI R2" + sheet KPI ROH/TOH/SQO Region)
import openpyxl, json, sys, re
src=sys.argv[1]; out='assets/js/kpi-data.js'
wb=openpyxl.load_workbook(src,data_only=True)
fnum=lambda v:float(v) if isinstance(v,(int,float)) else None
# --- rekap skor
ws=wb['KPI R2']
JAB=[('ROH','Regional Operation Head','REGIONAL OPERATION HEAD'),('TOH','Transaction Operation Head','TRANSACTION OPERATION HEAD'),
     ('SQOR','SQO Region','SQO REGION'),('ASM','Area Service Manager','AREA SERVICE MANAGER'),('SQOA','SQO Area','SQO AREA')]
cur=None;data={k:[] for k,_,_ in JAB}
for r in ws.iter_rows(min_row=4,values_only=True):
    a,b,c,d,e,f,g,h=r[:8]
    if isinstance(a,str) and not a.strip().isdigit():
        cur=next((k for k,_,p in JAB if a.upper().startswith(p)),None);continue
    if cur and c and fnum(h) is not None:
        data[cur].append({'rk':int(b),'n':str(c).strip(),'c':round(d,3),'i':round(e,3),'f':round(f,3),'d':round(g,3),'s':round(h,3),'r2':a is not None})
# --- rincian KPI Region II
def detail(sheet):
    w=wb[sheet];rows=[];pers=None;grp=None
    for r in w.iter_rows(min_row=7,values_only=True):
        a,b,c,d,e,f,g,h,i,j,k,l=(list(r)+[None]*12)[:12]
        if isinstance(b,str) and b.strip().isupper() and j is not None and f is None:
            pers=b.strip().title().replace('&','&');rows.append({'p':pers,'tot':round(j,3)});continue
        if f is not None and fnum(f) is not None and d:
            rows.append({'g':grp,'n':str(d).strip(),'w':fnum(c),'r':fnum(f),'t':fnum(g),'a':fnum(h),'v':fnum(j)});continue
        if d and f is None and j is None: grp=str(d).strip()
    return rows
det={'ROH':detail('KPI ROH'),'TOH':detail('KPI TOH'),'SQOR':detail('KPI SQO Region')}
# --- rincian KPI semua area (ASM, SQO Area): kolom realisasi/target/% diambil dari rumus SUMIF pada sheet KPI -> sheet Data
from openpyxl.utils import column_index_from_string as ci
wbf=openpyxl.load_workbook(src)
def area_detail(sheet,dsheet,codec,namec,r0,r1):
    wf,wv,dv=wbf[sheet],wb[sheet],wb[dsheet]
    meta=[];cols=[];grp=None
    for r in range(7,wv.max_row+1):
        b,c,d,f,j=[wv[f'{x}{r}'].value for x in 'BCDFJ']
        if isinstance(b,str) and b.strip().isupper() and j is not None and f is None:
            meta.append({'p':b.strip().title()});continue
        ff=wf[f'F{r}'].value
        if isinstance(ff,str) and 'SUMIF' in ff:
            cc=[re.search(r"'Data (?:ASM|SQO)'!([A-Z]+)\$\d+:",str(wf[f'{x}{r}'].value)).group(1) for x in 'FGH']
            jf=str(wf[f'J{r}'].value);th=float(re.search(r'H\d+>(\d+)%',jf).group(1))/100
            meta.append({'g':grp,'n':str(d).strip(),'w':round(float(c),5),'th':th,'fl':1 if '<0' in jf else 0});cols.append(cc)
        elif d and f is None and j is None: grp=str(d).strip()
    vals={};names={}
    for r in range(r0,r1+1):
        code=dv[f'{codec}{r}'].value;nm=dv[f'{namec}{r}'].value
        if code is None or nm is None or not str(nm).upper().startswith('AREA'):continue
        kode=str(int(code)) if isinstance(code,(int,float)) else str(code).strip();names[kode]=str(nm).strip()
        vals[kode]=[[ (round(v,5) if isinstance((v:=dv[f'{x}{r}'].value),(int,float)) else None) for x in cc] for cc in cols]
    return {'items':meta,'vals':vals,'names':names}
AD={'ASM':area_detail('KPI ASM','Data ASM','A','B',6,89),'SQOA':area_detail('KPI SQO Area','Data SQO','B','C',19,114)}
# pemetaan kode area: kode & nama diambil dari sheet Data ASM / Data SQO (nama kembar diberi kode sesuai urutan baris)
for k,ad in AD.items():
    by={}
    for kode,nm in ad['names'].items(): by.setdefault(nm,[]).append(kode)
    used={}
    for u in data[k]:
        ks=by.get(u['n'],[]);i=used.get(u['n'],0)
        if i<len(ks): u['k']=ks[i];used[u['n']]=i+1
    ad['dupes']={nm:ks for nm,ks in by.items() if len(ks)>1}
    bad=0
    for u in data[k]:
        v=ad['vals'].get(u.get('k'))
        if v is None: bad+=1;print('tanpa kode/rincian',k,u['n']);continue
        tot=0
        for m,(r_,t_,a_) in zip([m for m in ad['items'] if 'n' in m],v):
            a=a_ or 0;x=1.3*m['w']*100 if a>m['th'] else a*m['w']*100
            if m['fl'] and x<0:x=0
            tot+=x
        if abs(tot-u['s'])>0.02:bad+=1;print('selisih',k,u['k'],u['n'],round(tot,3),u['s'])
    print(k,'kode area',len(ad['vals']),'nama kembar',ad['dupes'],'selisih skor',bad)
J=[{'key':k,'nama':n,'unit':data[k]} for k,n,_ in JAB]
js='window.KITO_KPI='+json.dumps({'periode':'September 2026','fokus':'REGION II / SUMATERA 2','jabatan':J,'detail':det,'areaDetail':AD},ensure_ascii=False,separators=(',',':'))+';\n'
open(out,'w',encoding='utf-8').write(js)
print(len(js),{k:len(v) for k,v in data.items()},{k:len(v) for k,v in det.items()})
print([ (x['rk'],x['n']) for x in data['ROH'] if x['r2']],[ x['n'] for x in data['ASM'] if x['r2']],[x['n'] for x in data['SQOA'] if x['r2']])
for x in det['ROH'][:14]: print(x)
