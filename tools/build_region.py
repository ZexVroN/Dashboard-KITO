"""Bangun assets/js/region-data.js (Dashboard Region: 12 region + Nasional) dari tiga laporan bankwide.

Pemakaian (dari folder Dashboard-KITO):
    python tools/build_region.py --antrian <Dashboard_Antrian_Cabang.xlsx> \
        --o2o <Dashboard_Monitoring_Ganti_Kartu.xlsx> --nbds <Monitoring_Utilisasi_NBDS_sd_....xlsx>

Sumber:
    Waktu antri : sheet "Calc" (kunci R|<region>|<tgl> dan N|NASIONAL|<tgl>; "Gabungan" = MTD) -> P90 detik
    NBDS        : sheet "Region" (blok F P, F NP, NF, Agregat; baris per region + "Total" = Nasional)
                  MTD bulan berjalan = jumlah harian (konsisten dengan kolom bulanan laporan)
    O2O         : sheet "Region O2O" (segmen) dan "Regional by Channel" (Shifting Digital)
"""

import argparse
import datetime as dt
import json
import re

import openpyxl

JS = "assets/js/"
BULAN = {"januari": 1, "februari": 2, "maret": 3, "april": 4, "mei": 5, "juni": 6, "juli": 7,
         "agustus": 8, "september": 9, "oktober": 10, "november": 11, "desember": 12}
BULAN3 = {"Jan": 1, "Feb": 2, "Mar": 3, "Apr": 4, "May": 5, "Mei": 5, "Jun": 6, "Jul": 7, "Ags": 8, "Agu": 8,
          "Sep": 9, "Oct": 10, "Okt": 10, "Nov": 11, "Dec": 12, "Des": 12}
P90 = {"teller": 32, "cs": 37, "all": 27}          # kolom P90 di sheet Calc (0-based): AG, AL, AB
JML = {"teller": 28, "cs": 33, "all": 23}          # jumlah reservasi: AC, AH, X
NBDS_BLOK = [("(F P)", "fp"), ("(F NP)", "fnp"), ("(NF)", "nf"), ("Agregat", "agg")]


def detik(v):
    if v is None or isinstance(v, str):
        return None
    if isinstance(v, (dt.time, dt.datetime)):
        s = v.hour * 3600 + v.minute * 60 + v.second + v.microsecond / 1e6
    elif isinstance(v, dt.timedelta):
        s = v.total_seconds()
    elif isinstance(v, (int, float)):
        s = float(v) * 86400
    else:
        return None
    return round(s, 1)


def num(v):
    return int(round(v)) if isinstance(v, (int, float)) and not isinstance(v, bool) else 0


def label_nbds(lbl):
    if not isinstance(lbl, str):
        return None
    l = lbl.strip()
    m = re.fullmatch(r"(\d{1,2}) ([A-Za-z]+)(\d{2})", l)
    if m and m.group(2).lower() in BULAN:
        return f"20{m.group(3)}-{BULAN[m.group(2).lower()]:02d}-{int(m.group(1)):02d}"
    m = re.fullmatch(r"([A-Za-z]{3})(\d{2})", l)
    if m and m.group(1) in BULAN3:
        return f"M:20{m.group(2)}-{BULAN3[m.group(1)]:02d}"
    return None


def is_region(teks):
    """True untuk 'REGION II / SUMATERA 2' dan baris total nasional; bukan header 'Region'."""
    t = str(teks or "").strip()
    return bool(re.match(r"REGION\s+[IVXL]+\s*/", t, re.I)) or t.upper() in ("TOTAL", "GRAND TOTAL", "NASIONAL")


def ident(nama):
    """'REGION II / SUMATERA 2' -> ('II', 'Sumatera 2')"""
    m = re.match(r"REGION\s+([IVXL]+)\s*/\s*(.+)", str(nama).strip(), re.I)
    return (m.group(1).upper(), m.group(2).strip().title()) if m else (None, str(nama).strip().title())


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--antrian", required=True)
    ap.add_argument("--o2o", required=True)
    ap.add_argument("--nbds", required=True)
    a = ap.parse_args()
    reg = {}

    def ent(nama):
        key = str(nama).strip()
        if key.upper() in ("NASIONAL", "TOTAL", "GRAND TOTAL"):
            key = "NASIONAL"
        if key not in reg:
            no, nm = ident(key) if key != "NASIONAL" else (None, "Nasional")
            reg[key] = {"id": key, "no": no, "nama": nm, "waktu": {"d": {}, "m": None, "n": {}},
                        "nbds": {}, "o2o": {}, "ch": {}}
        return reg[key]

    # ----- 1. Waktu antri (Calc) -----
    wb = openpyxl.load_workbook(a.antrian, read_only=True, data_only=True)
    tgl_w = set()
    for r in wb["Calc"].iter_rows(min_row=2, values_only=True):
        k = r[2]
        if not isinstance(k, str):
            continue
        p = k.split("|")
        if p[0] not in ("R", "N"):
            continue
        e = ent(p[1])
        nilai = [detik(r[P90[f]]) for f in ("teller", "cs", "all")]
        jml = [num(r[JML[f]]) for f in ("teller", "cs", "all")]
        if p[2] == "Gabungan":
            e["waktu"]["m"], e["waktu"]["mn"] = nilai, jml
        else:
            d, m = p[2].split()
            iso = f"2026-{BULAN3[m]:02d}-{int(d):02d}"
            tgl_w.add(iso)
            e["waktu"]["d"][iso], e["waktu"]["n"][iso] = nilai, jml

    # ----- 2. NBDS (sheet Region) -----
    wb = openpyxl.load_workbook(a.nbds, read_only=True, data_only=True)
    blok, head = None, None
    tgl_n = set()
    for r in wb["Region"].iter_rows(values_only=True):
        b = r[1] if len(r) > 1 else None
        if isinstance(b, str):
            for tanda, k in NBDS_BLOK:
                if b.startswith("Frekuensi") and tanda in b:
                    blok, head = k, None
        if blok is None:
            continue
        lbl = [label_nbds(x) for x in r]
        if head is None and sum(1 for x in lbl if x and not x.startswith("M:")) > 10:
            head = lbl
            continue
        if head is None or not isinstance(b, str):
            continue
        if not is_region(b):
            continue
        e = ent(b)
        d = e["nbds"].setdefault(blok, {})
        for c, l in enumerate(head):
            if l and l not in d and c + 1 < len(r):
                d[l] = [num(r[c]), num(r[c + 1])]
                if not l.startswith("M:"):
                    tgl_n.add(l)
    akhir_n = max(tgl_n)
    bulan = akhir_n[:7]
    th, bl = map(int, bulan.split("-"))
    bulan_lalu = f"{th - (bl == 1)}-{(bl - 2) % 12 + 1:02d}"
    dates_n = sorted(t for t in tgl_n if t.startswith(bulan))
    for e in reg.values():
        for k, d in list(e["nbds"].items()):
            seri = [d.get(t) for t in dates_n]
            ada = [p for p in seri if p]
            e["nbds"][k] = {"d": seri,
                            "m": [sum(p[0] for p in ada), sum(p[1] for p in ada)] if ada else None,
                            "mp": d.get("M:" + bulan_lalu)}

    # ----- 3. O2O & Shifting -----
    wb = openpyxl.load_workbook(a.o2o, read_only=True, data_only=True)
    judul = " ".join(str(v) for v in next(wb["Region O2O"].iter_rows(min_row=2, max_row=2, values_only=True)) if v)
    mt = re.search(r"sd\s+(\d{1,2})\s+([A-Za-z]{3})\w*\s+(\d{2})", judul)
    tgl_o = f"20{mt.group(3)}-{BULAN3[mt.group(2).title()]:02d}-{int(mt.group(1)):02d}" if mt else akhir_n
    SEG = {"gold": 10, "platinum": 22, "silver": 26, "overall": 30}  # K, W, AA, AE (LIVIN; +1 NEW BDS)
    for r in wb["Region O2O"].iter_rows(min_row=5, max_row=40, values_only=True):
        if isinstance(r[1], str) and is_region(r[1]):
            ent(r[1])["o2o"] = {k: [num(r[c]), num(r[c + 1])] for k, c in SEG.items()}
    for r in wb["Regional by Channel"].iter_rows(min_row=5, max_row=40, values_only=True):
        if isinstance(r[1], str) and is_region(r[1]):
            ent(r[1])["ch"] = {"livin": num(r[2]), "eci": num(r[3]), "csm": num(r[4]), "o2o": num(r[5]),
                               "old": num(r[7]), "new": num(r[8])}

    nas = reg.pop("NASIONAL")
    urut = sorted(reg.values(), key=lambda e: ["I", "II", "III", "IV", "V", "VI", "VII", "VIII", "IX", "X", "XI", "XII"].index(e["no"]) if e["no"] else 99)
    out = {
        "tanggal": max(max(tgl_w), akhir_n), "bulan": bulan, "bulanLalu": bulan_lalu,
        "tglWaktu": sorted(tgl_w), "tglNbds": dates_n, "tglO2O": tgl_o,
        "fokus": "REGION II / SUMATERA 2",
        "sumber": {
            "waktu": re.sub(r".*[\\/]", "", a.antrian) + " · sheet Calc",
            "nbds": re.sub(r".*[\\/]", "", a.nbds) + " · sheet Region",
            "o2o": re.sub(r".*[\\/]", "", a.o2o) + " · sheet Region O2O & Regional by Channel",
        },
        "regions": urut, "nasional": nas,
    }
    with open(JS + "region-data.js", "w", encoding="utf-8") as f:
        f.write("window.KITO_REGION=" + json.dumps(out, ensure_ascii=False, separators=(",", ":")) + ";\n")
    lengkap = [e["id"] for e in urut if e["waktu"]["d"] and e["nbds"] and e["o2o"] and e["ch"]]
    print(f"region-data.js: {len(urut)} region (lengkap {len(lengkap)}) + Nasional · waktu {len(tgl_w)} tgl · "
          f"NBDS {len(dates_n)} tgl s.d. {akhir_n} · O2O per {tgl_o}")


if __name__ == "__main__":
    main()
