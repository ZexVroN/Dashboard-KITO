"""Konversi laporan "Dashboard Monitoring Utilisasi NBDS Region 2 DtD" (.xlsx)
menjadi assets/js/area-nbds-data.js untuk tampilan "Dashboard Utilisasi NBDS" di tab Area.

Pemakaian:
    python tools/convert_nbds.py <file.xlsx> [assets/js/area-nbds-data.js]

Sheet yang dibaca:
    - Area                      : New BDS / Old BDS per area (3 kanal), harian + MTD
    - Finansial Perorangan      : per cabang, Teller Perorangan (FP)
    - Finansial Non Perorangan  : per cabang, Teller Badan (FNP)
    - Non Finansial             : per cabang, Customer Service (NF)
    - Raw Data                  : baris transaksi Old BDS DtD (untuk rincian layanan, cabang, jenis transaksi)

Data pribadi (nama/no HP BOM-BOS) sengaja TIDAK diekspor; user ID disamarkan (******1234).
Hanya bulan laporan (harian + MTD) dan MTD bulan sebelumnya yang disimpan.
"""

import datetime as dt
import json
import re
import sys
from collections import defaultdict

import openpyxl

KANAL = [
    {"key": "fp", "label": "Teller Perorangan", "short": "Teller Perorangan"},
    {"key": "fnp", "label": "Teller Badan", "short": "Teller Badan"},
    {"key": "nf", "label": "Customer Service", "short": "CS"},
]
SHEET_CABANG = {
    "fp": "Finansial Perorangan",
    "fnp": "Finansial Non Perorangan",
    "nf": "Non Finansial",
}
# urutan seksi pada sheet "Area" (judul baris di atas header tanggal)
JUDUL_AREA = {
    "transaksi teller perorangan": "fp",
    "transaksi teller badan": "fnp",
    "transaksi customer service": "nf",
}
TARGET = 0.98


def num(v):
    """Angka sel -> int; teks/kosong/#DIV/0! -> None."""
    if isinstance(v, bool) or v is None:
        return None
    if isinstance(v, (int, float)):
        return int(round(v))
    return None


def peta_kolom(header):
    """Dari baris header, kembalikan [(label, kolom_newbds)].
    label = 'YYYY-MM-DD' untuk tanggal, 'MTD:YYYY-MM' untuk kolom MTD (bulan dari tanggal sebelumnya)."""
    hasil, bulan_terakhir = [], None
    for i, v in enumerate(header):
        if isinstance(v, dt.datetime):
            iso = v.date().isoformat()
            bulan_terakhir = iso[:7]
            hasil.append((iso, i))
        elif isinstance(v, str) and v.strip().upper().startswith("MTD") and bulan_terakhir:
            hasil.append(("MTD:" + bulan_terakhir, i))
    return hasil


def ambil_seri(row, kolom, bulan, bulan_lalu):
    """Pisahkan nilai [nb, ob] menjadi harian bulan laporan, MTD bulan laporan, dan MTD bulan lalu."""
    d, m, mp = {}, None, None
    for label, c in kolom:
        pair = [num(row[c]), num(row[c + 1])]
        if pair[0] is None and pair[1] is None:
            pair = None
        if label == "MTD:" + bulan:
            m = pair
        elif label == "MTD:" + bulan_lalu:
            mp = pair
        elif label.startswith(bulan):
            d[label] = pair
    return d, m, mp


def main(src, out):
    wb = openpyxl.load_workbook(src, data_only=True, read_only=False)

    # --- tanggal laporan = tanggal terakhir di header sheet Area ---
    ws = wb["Area"]
    rows = [list(r) for r in ws.iter_rows(values_only=True)]
    semua_tgl = sorted({v.date().isoformat() for r in rows for v in r if isinstance(v, dt.datetime)})
    tanggal = semua_tgl[-1]
    bulan = tanggal[:7]
    th, bl = map(int, bulan.split("-"))
    bulan_lalu = f"{th - (bl == 1)}-{(bl - 2) % 12 + 1:02d}"
    dates = [t for t in semua_tgl if t.startswith(bulan)]

    # --- sheet Area: 3 seksi, masing-masing: judul, header tanggal, sub-header, baris area ---
    areas = defaultdict(lambda: {"d": {}, "m": {}, "mp": {}})
    r = 0
    while r < len(rows):
        judul = str(rows[r][0] or "").strip().lower()
        if judul in JUDUL_AREA:
            k = JUDUL_AREA[judul]
            kolom = peta_kolom(rows[r + 1])
            r += 3  # lewati header tanggal + sub-header
            while r < len(rows) and rows[r][1]:
                nama = str(rows[r][1]).strip()
                d, m, mp = ambil_seri(rows[r], kolom, bulan, bulan_lalu)
                areas[nama]["d"][k] = [d.get(t) for t in dates]
                areas[nama]["m"][k] = m
                areas[nama]["mp"][k] = mp
                r += 1
        r += 1

    # --- sheet cabang ---
    cabang = {}
    for k, nama_sheet in SHEET_CABANG.items():
        rs = [list(x) for x in wb[nama_sheet].iter_rows(values_only=True)]
        hi = next(i for i, x in enumerate(rs) if str(x[1] or "").strip() == "Kode Cabang")
        kolom = peta_kolom(rs[hi])
        for x in rs[hi + 2:]:
            if num(x[1]) is None:
                continue
            kode = str(num(x[1]))
            c = cabang.setdefault(kode, {
                "kode": kode,
                "nama": str(x[2]).strip(),
                "area": str(x[3]).strip(),
                "kelas": str(x[4] or "").strip(),
                "d": {}, "m": {},
            })
            d, m, mp = ambil_seri(x, kolom, bulan, bulan_lalu)
            c["d"][k] = [d.get(t) for t in dates]
            c["m"][k] = m
            c.setdefault("mp", {})[k] = mp

    # --- angka area = jumlah cabang (sheet Area hanya untuk validasi) ---
    # Alasan: pada laporan 21 Okt 2026, Old BDS MTD Customer Service di sheet Area
    # tidak sama dengan jumlah cabang (bahkan melebihi jumlah harian sebulan).
    def jumlah(pairs):
        ada = [p for p in pairs if p]
        return [sum(p[0] or 0 for p in ada), sum(p[1] or 0 for p in ada)] if ada else None

    area_cabang = {}
    for c in cabang.values():
        a = area_cabang.setdefault(c["area"], {"d": {}, "m": {}, "mp": {}})
        for kn in KANAL:
            k = kn["key"]
            a["d"].setdefault(k, [[] for _ in dates])
            for i, p in enumerate(c["d"].get(k) or []):
                a["d"][k][i].append(p)
            a["m"].setdefault(k, []).append(c["m"].get(k))
            a["mp"].setdefault(k, []).append(c.get("mp", {}).get(k))
    for nama, a in area_cabang.items():
        for k in a["d"]:
            a["d"][k] = [jumlah(x) for x in a["d"][k]]
            a["m"][k] = jumlah(a["m"][k])
            a["mp"][k] = jumlah(a["mp"][k])
        lembar = areas.get(nama)
        if not lembar:
            print(f"PERINGATAN: area {nama} tidak ada di sheet Area")
            continue
        for k in a["d"]:
            # MTD bulan lalu di sheet Area hanya menjumlah tanggal yang tampil di sheet itu,
            # sedangkan kolom "MTD <bulan>" di sheet cabang lengkap; jadi hanya MTD berjalan yang divalidasi.
            if lembar["m"].get(k) and lembar["m"][k] != a["m"][k]:
                print(f"PERINGATAN: {nama} {k} MTD: sheet Area {lembar['m'][k]} ≠ Σ cabang {a['m'][k]} (dipakai Σ cabang)")
            beda = [dates[i] for i, (x, y) in enumerate(zip(lembar["d"].get(k, []), a["d"][k])) if x and x != y]
            if beda:
                print(f"PERINGATAN: {nama} {k} harian beda pada {beda}")
    areas = area_cabang
    for c in cabang.values():
        c.pop("mp", None)  # MTD bulan lalu per cabang tidak dipakai di tampilan

    # --- Raw Data: baris transaksi yang masih lewat Old BDS (DtD tanggal laporan) ---
    # Format baris: [kode, area, kanal, fitur, jenis transaksi, nasabah, user ID tersamar, total]
    # User ID disamarkan seperti data Old BDS bawaan (data.js): "******" + 4 digit terakhir.
    oldbds = []
    rs = wb["Raw Data"].iter_rows(values_only=True)
    head = [str(h or "").strip() for h in next(rs)]
    ix = {h: i for i, h in enumerate(head)}
    tgl_raw = int(tanggal.replace("-", ""))
    teks = lambda v: str(v if v is not None else "").strip()
    for x in rs:
        if x[ix["flag_bds"]] != "OLD BDS" or num(x[ix["time_pr"]]) != tgl_raw:
            continue
        kat = x[ix["Kategorisasi"]]
        k = "nf" if kat == "Customer Service" else ("fp" if x[ix["cif_class"]] == "perorangan" else "fnp")
        uid = teks(x[ix["user_id"]])
        oldbds.append([
            teks(num(x[ix["kode_cabang"]]) if num(x[ix["kode_cabang"]]) is not None else x[ix["kode_cabang"]]).lstrip("0") or "0",
            teks(x[ix["area"]]),
            k,
            teks(x[ix["fitur"]]),
            teks(x[ix["service"]]),
            teks(x[ix["cif_class"]]),
            ("******" + uid[-4:]) if uid else "",
            num(x[ix["tot_frek"]]) or 0,
        ])
    oldbds.sort(key=lambda r: (r[1], r[0], r[2], r[3], -r[7]))
    tak_dikenal = sorted({r[0] for r in oldbds} - set(cabang))
    if tak_dikenal:
        print(f"PERINGATAN: kode cabang Raw Data tidak ada di sheet cabang: {tak_dikenal}")

    data = {
        "tanggal": tanggal,
        "bulan": bulan,
        "bulanLalu": bulan_lalu,
        "sumber": re.sub(r".*[\\/]", "", src) + " · sheet Area, Finansial Perorangan, Finansial Non Perorangan, Non Finansial, Raw Data",
        "target": TARGET,
        "kanal": KANAL,
        "dates": dates,
        "areas": [{"name": n, **v} for n, v in areas.items()],
        "cabang": list(cabang.values()),
        "oldbds": oldbds,
    }
    teks = "window.KITO_AREA_NBDS=" + json.dumps(data, ensure_ascii=False, separators=(",", ":")) + ";\n"
    with open(out, "w", encoding="utf-8") as fh:
        fh.write(teks)
    print(f"{out}: {len(data['areas'])} area, {len(data['cabang'])} cabang, {len(dates)} tanggal ({dates[0]}..{dates[-1]}), {len(teks)/1024:.1f} KB")


if __name__ == "__main__":
    if len(sys.argv) < 2:
        sys.exit(__doc__)
    main(sys.argv[1], sys.argv[2] if len(sys.argv) > 2 else "assets/js/area-nbds-data.js")
