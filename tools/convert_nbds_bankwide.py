"""Konversi laporan bankwide "Monitoring Utilisasi NBDS sd <tgl>" (.xlsx) untuk Region II.

Pemakaian (dari folder Dashboard-KITO):
    python tools/convert_nbds_bankwide.py <Monitoring_Utilisasi_NBDS_sd_....xlsx>

Sheet yang dibaca:
    Cab F P / Cab F NP / Cab NF : New BDS & Old BDS harian per cabang (Teller Perorangan, Teller Badan, CS)
    Raw                         : baris transaksi tanggal laporan -> rincian Old BDS
Hasil:
    assets/js/area-nbds-data.js : dashboard Utilisasi NBDS (format sama dengan convert_nbds.py)
    assets/js/data.js           : D.M[i][4..9] (utilisasi DtD tanggal laporan) dan D.O (detail Old BDS)

Catatan metode:
    - MTD bulan berjalan = jumlah harian tanggal 1 s.d. tanggal laporan. Ini konsisten dengan kolom
      bulanan laporan ini (mis. "Sep26" = jumlah harian September). Kolom "1-19Okt26" tidak dipakai
      karena hanya sampai tanggal 19.
    - MTD bulan lalu = kolom bulanan laporan (mis. "Sep26").
    - Nama/no HP BOM-BOS tidak diekspor; user ID disamarkan (******1234).
"""

import datetime as dt
import json
import re
import sys

import openpyxl

R2 = "REGION II / SUMATERA 2"
JS = "assets/js/"
BULAN = {"januari": 1, "februari": 2, "maret": 3, "april": 4, "mei": 5, "juni": 6, "juli": 7,
         "agustus": 8, "september": 9, "oktober": 10, "november": 11, "desember": 12}
BULAN3 = {"Jan": 1, "Feb": 2, "Mar": 3, "Apr": 4, "May": 5, "Jun": 6, "Jul": 7, "Ags": 8, "Sep": 9,
          "Oct": 10, "Okt": 10, "Nov": 11, "Dec": 12, "Des": 12}
SHEET = {"fp": "Cab F P", "fnp": "Cab F NP", "nf": "Cab NF"}
KANAL = [
    {"key": "fp", "label": "Teller Perorangan", "short": "Teller Perorangan"},
    {"key": "fnp", "label": "Teller Badan", "short": "Teller Badan"},
    {"key": "nf", "label": "Customer Service", "short": "CS"},
]
TARGET = 0.98


def num(v):
    return int(round(v)) if isinstance(v, (int, float)) and not isinstance(v, bool) else None


def baca_js(path, awalan):
    teks = open(path, encoding="utf-8").read().strip()
    assert teks.startswith(awalan) and teks.endswith(";"), path
    return json.loads(teks[len(awalan):-1])


def tulis_js(path, awalan, data):
    with open(path, "w", encoding="utf-8") as f:
        f.write(awalan + json.dumps(data, ensure_ascii=False, separators=(",", ":")) + ";\n")


def label_kolom(lbl, tahun):
    """'22 Oktober26' -> '2026-10-22'; 'Sep26' -> 'M:2026-09'; selain itu None."""
    if not isinstance(lbl, str):
        return None
    l = lbl.strip()
    m = re.fullmatch(r"(\d{1,2}) ([A-Za-z]+)(\d{2})", l)
    if m and m.group(2).lower() in BULAN:
        return f"20{m.group(3)}-{BULAN[m.group(2).lower()]:02d}-{int(m.group(1)):02d}"
    m = re.fullmatch(r"([A-Za-z]{3})(\d{2})", l)
    if m and m.group(1) in BULAN3:
        return f"M:20{m.group(2)}-{BULAN3[m.group(1)]:02d}"
    return None


def jumlah(pairs):
    ada = [p for p in pairs if p]
    return [sum(p[0] or 0 for p in ada), sum(p[1] or 0 for p in ada)] if ada else None


def main(src):
    D = baca_js(JS + "data.js", "const D = ")
    N_lama = baca_js(JS + "area-nbds-data.js", "window.KITO_AREA_NBDS=")
    nama_lama = {c["kode"]: c["nama"] for c in N_lama["cabang"]}
    kode_ke_i = {str(b[0]): i for i, b in enumerate(D["B"])}

    wb = openpyxl.load_workbook(src, read_only=True, data_only=True)
    data = {}          # kode -> kanal -> {label: [nb, ob]}
    nama_sheet = {}
    for k, sh in SHEET.items():
        head = None
        for i, r in enumerate(wb[sh].iter_rows(values_only=True), 1):
            if i == 6:
                head = [label_kolom(x, 2026) for x in r]
                continue
            if i < 8 or r[3] != R2:
                continue
            kode = str(num(r[1]))
            nama_sheet[kode] = str(r[2]).strip()
            d = data.setdefault(kode, {}).setdefault(k, {})
            for c, lbl in enumerate(head):
                # kolom "Growth" di ujung kanan mengulang label tanggal: pakai kemunculan pertama
                if lbl and lbl not in d and c + 1 < len(r):
                    d[lbl] = [num(r[c]), num(r[c + 1])]

    harian = sorted({t for x in data.values() for d in x.values() for t in d if not t.startswith("M:")})
    tanggal = harian[-1]
    bulan = tanggal[:7]
    th, bl = map(int, bulan.split("-"))
    bulan_lalu = f"{th - (bl == 1)}-{(bl - 2) % 12 + 1:02d}"
    dates = [t for t in harian if t.startswith(bulan)]

    cabang = []
    for b in D["B"]:
        kode = str(b[0])
        x = data.get(kode)
        if not x:
            print("PERINGATAN: cabang tidak ada di laporan NBDS:", kode, b[1])
            continue
        c = {"kode": kode, "nama": nama_lama.get(kode, nama_sheet.get(kode, b[1])), "area": b[2],
             "kelas": b[3], "d": {}, "m": {}, "mp": {}}
        for kn in KANAL:
            k = kn["key"]
            seri = [x.get(k, {}).get(t) for t in dates]
            c["d"][k] = seri
            c["m"][k] = jumlah(seri)
            c["mp"][k] = x.get(k, {}).get("M:" + bulan_lalu)
        cabang.append(c)

    areas = {}
    for c in cabang:
        a = areas.setdefault(c["area"], {"name": c["area"], "d": {}, "m": {}, "mp": {}, "_c": []})
        a["_c"].append(c)
    for a in areas.values():
        for kn in KANAL:
            k = kn["key"]
            a["d"][k] = [jumlah([c["d"][k][i] for c in a["_c"]]) for i in range(len(dates))]
            a["m"][k] = jumlah([c["m"][k] for c in a["_c"]])
            a["mp"][k] = jumlah([c["mp"][k] for c in a["_c"]])
        del a["_c"]
    for c in cabang:
        del c["mp"]

    # ----- Raw: baris Old BDS tanggal laporan -----
    teks = lambda v: str(v if v is not None else "").strip()
    it = wb["Raw"].iter_rows(values_only=True)
    ix = {str(h).strip(): i for i, h in enumerate(next(it)) if h}
    tgl_raw = int(tanggal.replace("-", ""))
    oldbds, D_O = [], []
    for r in it:
        # time_pr dan kode_cabang bisa tersimpan sebagai angka atau teks
        if r[ix["flag_bds"]] != "OLD BDS" or teks(r[ix["time_pr"]]).split(".")[0] != str(tgl_raw):
            continue
        kode = teks(r[ix["kode_cabang"]]).split(".")[0].lstrip("0")
        if kode not in kode_ke_i:
            continue
        kat = teks(r[ix["Kategorisasi"]])
        cif = teks(r[ix["cif_class"]])
        k = "nf" if kat.lower().startswith("non") else ("fp" if cif == "perorangan" else "fnp")
        uid = teks(r[ix["user_id"]])
        uid = ("******" + uid[-4:]) if uid else ""
        frek = num(r[ix["tot_frek"]]) or 0
        i = kode_ke_i[kode]
        oldbds.append([kode, D["B"][i][2], k, teks(r[ix["fitur"]]), teks(r[ix["service"]]), cif, uid, frek])
        D_O.append([i, teks(r[ix["service"]]), cif, uid, "Customer Service" if k == "nf" else "Teller", "OBDS", frek])
    oldbds.sort(key=lambda r: (r[1], r[0], r[2], r[3], -r[7]))
    D_O.sort(key=lambda r: (r[0], r[1]))

    # validasi: total Old BDS Raw per cabang & kanal = Old BDS harian di sheet cabang
    j = dates.index(tanggal)
    beda = []
    for c in cabang:
        for kn in KANAL:
            k = kn["key"]
            raw = sum(r[7] for r in oldbds if r[0] == c["kode"] and r[2] == k)
            sh = (c["d"][k][j] or [0, 0])[1] or 0
            if raw != sh:
                beda.append((c["kode"], k, raw, sh))
    if beda:
        print(f"PERINGATAN: {len(beda)} sel Old BDS Raw ≠ sheet cabang, contoh {beda[:5]}")

    N = {
        "tanggal": tanggal, "bulan": bulan, "bulanLalu": bulan_lalu,
        "sumber": re.sub(r".*[\\/]", "", src) + " · sheet Cab F P, Cab F NP, Cab NF, Raw",
        "target": TARGET, "kanal": KANAL, "dates": dates,
        "areas": list(areas.values()), "cabang": cabang, "oldbds": oldbds,
        "metodeMTD": "Jumlah harian 1 s.d. tanggal laporan (konsisten dengan kolom bulanan laporan)",
    }
    tulis_js(JS + "area-nbds-data.js", "window.KITO_AREA_NBDS=", N)

    # data.js: D.M[4..9] utilisasi DtD tanggal laporan + D.O
    by = {c["kode"]: c for c in cabang}
    for kode, i in kode_ke_i.items():
        m = D["M"][i]
        c = by.get(kode)
        if m is None or not c:
            continue
        vals = []
        for k in ("fp", "fnp", "nf"):
            p = c["d"][k][j] or [0, 0]
            vals += [p[0] or 0, p[1] or 0]
        m[4:10] = vals
    D["O"] = D_O
    tulis_js(JS + "data.js", "const D = ", D)
    print(f"area-nbds-data.js: {len(cabang)} cabang, {len(areas)} area, {len(dates)} tanggal ({dates[0]}..{dates[-1]}), "
          f"Old BDS {len(oldbds)} baris / {sum(r[7] for r in oldbds)} trx · data.js: D.M[4..9] & D.O per {tanggal}")


if __name__ == "__main__":
    if len(sys.argv) < 2:
        sys.exit(__doc__)
    main(sys.argv[1])
