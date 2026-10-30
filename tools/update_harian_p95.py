"""Update harian dashboard KITO dari 3 laporan bankwide, untuk laporan antrian format P95 (satu hari).

Pemakaian (dari folder Dashboard-KITO):
    python tools/update_harian_p95.py --antrian <DDMMYYYY_Dashboard_Antrian_Cabang.xlsx> \
        --o2o <Dashboard_Monitoring_Ganti_Kartu_DD_Okt_26.xlsx> \
        --nbds <Monitoring_Utilisasi_NBDS_sd_DD_Oktober_2026.xlsx>

Beda dengan update_semua.py: file antrian ini tidak punya sheet "Calc" (riwayat P90 banyak tanggal),
melainkan sheet Bankwide / Region / Area / Cabang berisi P95 untuk SATU tanggal (isi sheet RAW).
Jadi tanggal itu DITAMBAHKAN ke riwayat yang sudah ada; riwayat P90 dan MTD tidak disentuh.

Yang diperbarui:
  data.js            : dates, Q (Teller, CS, All per cabang), queueMetricByDate, M[0..3,10] (O2O, CSM, Old BDS,
                       NBDS kanal), bw, mdate; M[4..9] dan D.O lewat convert_nbds_bankwide.py
  area-waktu-data.js : area + Region II untuk tanggal baru (Teller, CS, All), metricByDate
  area-kelompok-data : DtD per cabang = P95 tanggal baru (MTD tetap)
  area-o2o-data.js   : utilisasi O2O per segmen (Area O2O, Region O2O)
  area-nbds-data.js  : lewat convert_nbds_bankwide.py
  region-data.js     : waktu antri (12 region + nasional) tanggal baru, NBDS (sheet Region), O2O & Shifting
"""

import argparse
import datetime as dt
import json
import os
import re
import subprocess
import sys

import openpyxl

R2 = "REGION II / SUMATERA 2"
JS = "assets/js/"
BULAN = {"Jan": 1, "Feb": 2, "Mar": 3, "Apr": 4, "Mei": 5, "May": 5, "Jun": 6, "Jul": 7, "Agu": 8, "Agt": 8,
         "Ags": 8, "Sep": 9, "Okt": 10, "Oct": 10, "Nov": 11, "Des": 12, "Dec": 12}
BULAN_PANJANG = {"januari": 1, "februari": 2, "maret": 3, "april": 4, "mei": 5, "juni": 6, "juli": 7,
                 "agustus": 8, "september": 9, "oktober": 10, "november": 11, "desember": 12}
NBDS_BLOK = [("(F P)", "fp"), ("(F NP)", "fnp"), ("(NF)", "nf"), ("Agregat", "agg")]
URUT_ROMAWI = ["I", "II", "III", "IV", "V", "VI", "VII", "VIII", "IX", "X", "XI", "XII"]


# ---------- util ----------
def baca_js(path, awalan):
    teks = open(path, encoding="utf-8").read()
    isi = teks.strip()
    assert isi.startswith(awalan) and isi.endswith(";"), path
    return json.loads(isi[len(awalan):-1]), teks.endswith("\n")


def tulis_js(path, awalan, data, newline):
    with open(path, "w", encoding="utf-8", newline="") as f:
        f.write(awalan + json.dumps(data, ensure_ascii=False, separators=(",", ":")) + ";" + ("\n" if newline else ""))


def detik(v):
    if v is None or isinstance(v, str):
        return None
    if isinstance(v, (dt.time, dt.datetime)):
        s = v.hour * 3600 + v.minute * 60 + v.second + v.microsecond / 1e6
    elif isinstance(v, dt.timedelta):
        s = v.total_seconds()
    elif isinstance(v, (int, float)):
        s = float(v) * 86400
    else:
        return None
    return round(s, 1)


def angka(v):
    return int(round(v)) if isinstance(v, (int, float)) and not isinstance(v, bool) else 0


def num(v):
    return int(round(v)) if isinstance(v, (int, float)) and not isinstance(v, bool) else None


def norm(s):
    return re.sub(r"[^a-z0-9]", "", re.sub(r"^area\s+", "", str(s or "").lower()))


def p95(nilai, jumlah):
    """P95 dalam detik; None bila tidak ada antrian untuk layanan itu."""
    if not jumlah:
        return None
    return detik(nilai)


def tgl_judul(teks, default=None):
    m = re.search(r"sd\s+(\d{1,2})\s+([A-Za-z]{3})\w*\s+(\d{2})", teks)
    return f"20{m.group(3)}-{BULAN[m.group(2).title()]:02d}-{int(m.group(1)):02d}" if m else default


def label_nbds(lbl):
    if not isinstance(lbl, str):
        return None
    l = lbl.strip()
    m = re.fullmatch(r"(\d{1,2}) ([A-Za-z]+)(\d{2})", l)
    if m and m.group(2).lower() in BULAN_PANJANG:
        return f"20{m.group(3)}-{BULAN_PANJANG[m.group(2).lower()]:02d}-{int(m.group(1)):02d}"
    m = re.fullmatch(r"([A-Za-z]{3})(\d{2})", l)
    if m and m.group(1) in BULAN:
        return f"M:20{m.group(2)}-{BULAN[m.group(1)]:02d}"
    return None


def is_region(teks):
    t = str(teks or "").strip()
    return bool(re.match(r"REGION\s+[IVXL]+\s*/", t, re.I)) or t.upper() in ("TOTAL", "GRAND TOTAL", "NASIONAL")


def ident(nama):
    m = re.match(r"REGION\s+([IVXL]+)\s*/\s*(.+)", str(nama).strip(), re.I)
    return (m.group(1).upper(), m.group(2).strip().title()) if m else (None, str(nama).strip().title())


def nama_file(p):
    return re.sub(r".*[\\/]", "", p)


# ---------- 1. antrian P95 ----------
def baca_antrian(path):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    baris1 = next(wb["RAW"].iter_rows(min_row=2, max_row=2, values_only=True))
    v = baris1[0]
    if isinstance(v, (dt.datetime, dt.date)):
        tgl = v.date().isoformat() if isinstance(v, dt.datetime) else v.isoformat()
    else:
        tgl = str(v)[:10]
    m = re.match(r"(\d{2})(\d{2})(\d{4})", nama_file(path))
    if m and f"{m.group(3)}-{m.group(2)}-{m.group(1)}" != tgl:
        sys.exit(f"Tanggal RAW ({tgl}) tidak sama dengan tanggal pada nama file ({nama_file(path)}).")

    cab = {}
    for r in wb["Cabang"].iter_rows(min_row=6, values_only=True):
        if r[1] == R2 and r[4] is not None:
            cab[str(r[4])] = {"n": [angka(r[10]), angka(r[13]), angka(r[7])],
                              "v": [p95(r[11], r[10]), p95(r[14], r[13]), p95(r[8], r[7])]}  # teller, cs, all
    area = {}
    for r in wb["Area"].iter_rows(min_row=6, values_only=True):
        if r[1] == R2 and r[3]:
            area[norm(r[3])] = [p95(r[10], r[9]), p95(r[13], r[12]), p95(r[7], r[6])]
    reg = {}
    for r in wb["Region"].iter_rows(min_row=6, values_only=True):
        if isinstance(r[3], str) and is_region(r[3]):
            reg[r[3].strip()] = {"v": [p95(r[10], r[9]), p95(r[13], r[12]), p95(r[7], r[6])],
                                 "n": [angka(r[9]), angka(r[12]), angka(r[6])]}
    for r in wb["Bankwide"].iter_rows(min_row=6, max_row=6, values_only=True):
        reg["NASIONAL"] = {"v": [p95(r[10], r[9]), p95(r[13], r[12]), p95(r[7], r[6])],
                           "n": [angka(r[9]), angka(r[12]), angka(r[6])]}
    return tgl, cab, area, reg


# ---------- 2. O2O ----------
def baca_o2o(path):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    o2o, ch = {}, {}
    for r in wb["Cabang O2O"].iter_rows(min_row=5, values_only=True):
        if r[2] == R2:
            o2o[str(r[1])] = (angka(r[35]), angka(r[36]))           # Overall O2O: Livin, New BDS
    for r in wb["Cabang by Channel"].iter_rows(min_row=5, values_only=True):
        if r[2] == R2:
            ch[str(r[1])] = {"csm": angka(r[9]), "old": angka(r[12]), "new": angka(r[13])}
    judul = " ".join(str(v) for v in next(wb["Cabang O2O"].iter_rows(min_row=1, max_row=1, values_only=True)) if v)
    tgl = tgl_judul(judul)
    if not tgl:
        sys.exit("Tanggal laporan O2O tidak ditemukan pada judul sheet Cabang O2O.")

    A_COL = {"gold": 11, "platinum": 23, "silver": 27, "overall": 31}
    R_COL = {"gold": 10, "platinum": 22, "silver": 26, "overall": 30}

    def segmen(row, cols, banding=None):
        out = {}
        for k, c in cols.items():
            lv, nb = angka(row[c]), angka(row[c + 1])
            t = lv + nb
            p = lv / t if t else 0
            ket = "avg" if banding is None else ("above" if p >= banding[k][3] else "below")
            out[k] = [lv, nb, t, p, ket]
        return out

    bank = reg2_row = None
    bw = None
    for r in wb["Region O2O"].iter_rows(min_row=5, max_row=40, values_only=True):
        if isinstance(r[1], str) and r[1].strip().upper() == "GRAND TOTAL":
            bank = segmen(r, R_COL)
            lv, nb = angka(r[30]), angka(r[31])
            bw = lv / (lv + nb) if lv + nb else None
        elif r[1] == R2:
            reg2_row = r
    area_rows = [r for r in wb["Area O2O "].iter_rows(min_row=4, values_only=True) if r[1] == R2 and r[2]]
    reg2 = segmen(reg2_row, R_COL, bank)
    areas = [{"name": str(r[2]).strip(), "rank": None, "s": segmen(r, A_COL, reg2)} for r in area_rows]
    areas.sort(key=lambda a: -a["s"]["overall"][3])
    for i, a in enumerate(areas, 1):
        a["rank"] = i

    # Region O2O (segmen) + Regional by Channel (Shifting Digital) untuk region-data.js
    SEG = {"gold": 10, "platinum": 22, "silver": 26, "overall": 30}
    reg_o2o, reg_ch = {}, {}
    for r in wb["Region O2O"].iter_rows(min_row=5, max_row=40, values_only=True):
        if isinstance(r[1], str) and is_region(r[1]):
            reg_o2o[r[1].strip()] = {k: [num(r[c]), num(r[c + 1])] for k, c in SEG.items()}
    for r in wb["Regional by Channel"].iter_rows(min_row=5, max_row=40, values_only=True):
        if isinstance(r[1], str) and is_region(r[1]):
            reg_ch[r[1].strip()] = {"livin": num(r[2]), "eci": num(r[3]), "csm": num(r[4]), "o2o": num(r[5]),
                                    "old": num(r[7]), "new": num(r[8])}
    return {"tgl": tgl, "o2o": o2o, "ch": ch, "bw": bw, "bank": bank, "reg2": reg2, "areas": areas,
            "reg_o2o": reg_o2o, "reg_ch": reg_ch}


# ---------- 3. NBDS untuk region-data ----------
def baca_nbds_region(path):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    blok = head = None
    data, tgl_n = {}, set()
    for r in wb["Region"].iter_rows(values_only=True):
        b = r[1] if len(r) > 1 else None
        if isinstance(b, str):
            for tanda, k in NBDS_BLOK:
                if b.startswith("Frekuensi") and tanda in b:
                    blok, head = k, None
        if blok is None:
            continue
        lbl = [label_nbds(x) for x in r]
        if head is None and sum(1 for x in lbl if x and not x.startswith("M:")) > 10:
            head = lbl
            continue
        if head is None or not isinstance(b, str) or not is_region(b):
            continue
        key = b.strip()
        if key.upper() in ("TOTAL", "GRAND TOTAL", "NASIONAL"):
            key = "NASIONAL"
        d = data.setdefault(key, {}).setdefault(blok, {})
        for c, l in enumerate(head):
            if l and l not in d and c + 1 < len(r):
                d[l] = [num(r[c]), num(r[c + 1])]
                if not l.startswith("M:"):
                    tgl_n.add(l)
    akhir = max(tgl_n)
    bulan = akhir[:7]
    th, bl = map(int, bulan.split("-"))
    lalu = f"{th - (bl == 1)}-{(bl - 2) % 12 + 1:02d}"
    tanggal = sorted(t for t in tgl_n if t.startswith(bulan))
    hasil = {}
    for key, blokk in data.items():
        hasil[key] = {}
        for k, d in blokk.items():
            seri = [d.get(t) for t in tanggal]
            ada = [p for p in seri if p]
            hasil[key][k] = {"d": seri, "m": [sum(p[0] for p in ada), sum(p[1] for p in ada)] if ada else None,
                             "mp": d.get("M:" + lalu)}
    return hasil, tanggal, bulan, lalu


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--antrian", required=True)
    ap.add_argument("--o2o", required=True)
    ap.add_argument("--nbds", required=True)
    ap.add_argument("--lewati-nbds-area", action="store_true", help="jangan jalankan convert_nbds_bankwide.py")
    a = ap.parse_args()

    # NBDS (area-nbds-data.js, D.M[4..9], D.O) memakai skrip yang sudah ada; harus lebih dulu karena menulis data.js
    if not a.lewati_nbds_area:
        subprocess.run([sys.executable, os.path.join("tools", "convert_nbds_bankwide.py"), a.nbds], check=True)

    tgl, cab, area_q, reg_q = baca_antrian(a.antrian)
    O = baca_o2o(a.o2o)
    nbds_reg, tgl_nbds, bulan, bulan_lalu = baca_nbds_region(a.nbds)

    # ===== data.js =====
    D, nl_d = baca_js(JS + "data.js", "const D = ")
    kode_ke_i = {str(b[0]): i for i, b in enumerate(D["B"])}
    kurang = [k for k in kode_ke_i if k not in cab]
    if kurang:
        print("PERINGATAN: cabang tidak ada di laporan antrian:", kurang)
    if tgl not in D["dates"]:
        D["dates"].append(tgl)
        D["dates"].sort()
    j = D["dates"].index(tgl)
    for i, b in enumerate(D["B"]):
        v = cab.get(str(b[0]), {}).get("v", [None, None, None])
        seri = D["Q"][i]
        while len(seri) < len(D["dates"]):
            seri.append(None)
        seri[j] = v
        D["Q"][i] = seri
    D.setdefault("queueMetricByDate", {})[tgl] = "P95"

    beda_m = 0
    for kode, i in kode_ke_i.items():
        m = D["M"][i]
        if m is None:
            continue
        lama = list(m[:4]) + [m[10] if len(m) > 10 else None]
        if kode in O["o2o"]:
            m[0], m[1] = O["o2o"][kode]
        if kode in O["ch"]:
            m[2], m[3] = O["ch"][kode]["csm"], O["ch"][kode]["old"]
            while len(m) < 11:
                m.append(0)
            m[10] = O["ch"][kode]["new"]
        beda_m += lama != list(m[:4]) + [m[10] if len(m) > 10 else None]
    if O["bw"] is not None:
        D["bw"] = O["bw"]
    D["mdate"] = O["tgl"]
    tulis_js(JS + "data.js", "const D = ", D, nl_d)

    # ===== area-waktu-data.js =====
    W, nl_w = baca_js(JS + "area-waktu-data.js", "window.KITO_AREA_WAKTU=")

    def baris(v):  # urutan fields: T single, T multipel, Teller, CS single, CS multipel, CS, All Trx
        return [None, None, v[0], None, None, v[1], v[2]]
    for ent in W["areas"]:
        v = area_q.get(norm(ent["name"]))
        if v is None:
            print("PERINGATAN: area tidak ditemukan di laporan antrian:", ent["name"])
            continue
        ent["d"][tgl] = baris(v)
    r2v = reg_q[R2]["v"]
    W["region"]["d"][tgl] = baris(r2v)
    if tgl not in W["dates"]:
        W["dates"].append(tgl)
        W["dates"].sort()
    W["tanggal"] = tgl
    W.setdefault("metricByDate", {})[tgl] = "P95"
    W["sumber"] = f"{nama_file(a.antrian)} · P95 {tgl}, historis P90"

    # ===== area-kelompok-data.js: DtD = P95 tanggal baru =====
    K, nl_k = baca_js(JS + "area-kelompok-data.js", "window.KITO_AREA_KELOMPOK=")
    for r in K:
        v = cab.get(str(r["kode"]), {}).get("v")
        if v:
            r["dtd"] = {"teller": v[0], "cs": v[1], "all": v[2]}
    tulis_js(JS + "area-kelompok-data.js", "window.KITO_AREA_KELOMPOK=", K, nl_k)

    # ===== area-o2o-data.js =====
    AO, nl_o = baca_js(JS + "area-o2o-data.js", "window.KITO_AREA_O2O=")
    AO["region"]["s"], AO["bankwide"]["s"], AO["areas"] = O["reg2"], O["bank"], O["areas"]
    AO["tanggal"] = O["tgl"]
    AO["sumber"] = nama_file(a.o2o) + " · sheet Area O2O & Region O2O"
    tulis_js(JS + "area-o2o-data.js", "window.KITO_AREA_O2O=", AO, nl_o)

    # ===== region-data.js =====
    R, nl_r = baca_js(JS + "region-data.js", "window.KITO_REGION=")
    semua = {e["id"]: e for e in R["regions"]}
    semua["NASIONAL"] = R["nasional"]
    for key, e in semua.items():
        q = reg_q.get(key)
        if q:
            e["waktu"]["d"][tgl] = q["v"]
            e["waktu"].setdefault("n", {})[tgl] = q["n"]
        else:
            print("PERINGATAN: region tidak ditemukan di laporan antrian:", key)
        if key in nbds_reg:
            e["nbds"] = nbds_reg[key]
        if key in O["reg_o2o"]:
            e["o2o"] = O["reg_o2o"][key]
        if key in O["reg_ch"]:
            e["ch"] = O["reg_ch"][key]
    R["tglWaktu"] = sorted(set(R["tglWaktu"]) | {tgl})
    R["tglNbds"], R["bulan"], R["bulanLalu"] = tgl_nbds, bulan, bulan_lalu
    R["tglO2O"] = O["tgl"]
    R["tanggal"] = max(tgl, tgl_nbds[-1])
    R.setdefault("queueMetricByDate", {})[tgl] = "P95"
    R["sumber"] = {
        "waktu": f"{nama_file(a.antrian)} · DtD {tgl} P95 seluruh 12 Region; MTD dan riwayat sampai 26 Okt P90",
        "nbds": f"{nama_file(a.nbds)} · sheet Region · DtD dan MTD s.d. {tgl_nbds[-1]} · 12 Region + Bankwide",
        "o2o": nama_file(a.o2o),
    }
    tulis_js(JS + "region-data.js", "window.KITO_REGION=", R, nl_r)

    # sinkronkan Region II di area-waktu dengan region-data (sumber sama: sheet Region laporan antrian)
    r2e = semua[R2]["waktu"]["d"]
    diperbaiki = []
    for t, v in r2e.items():
        cur = W["region"]["d"].get(t)
        if cur and [cur[2], cur[5], cur[6]] != [v[0], v[1], v[2]]:
            diperbaiki.append((t, [cur[2], cur[5], cur[6]], v))
            W["region"]["d"][t] = baris(v)
    tulis_js(JS + "area-waktu-data.js", "window.KITO_AREA_WAKTU=", W, nl_w)

    print(f"data.js            : tanggal {tgl} ditambah (P95) · M O2O/CSM/OldBDS/NBDS-kanal berubah {beda_m} cabang · bw {D['bw']:.4%} · mdate {D['mdate']}")
    print(f"area-waktu-data.js : {len(W['areas'])} area + Region II, tanggal {W['tanggal']}")
    if diperbaiki:
        for t, lama, baru in diperbaiki:
            print(f"  Region II {t} disinkronkan dengan region-data: [Teller, CS, All] {lama} -> {baru}")
    print(f"area-kelompok-data : DtD = P95 {tgl} untuk {len(K)} cabang")
    print(f"area-o2o-data.js   : {len(AO['areas'])} area, tanggal {AO['tanggal']}")
    print(f"region-data.js     : {len(R['regions'])} region + nasional · waktu {tgl} · NBDS s.d. {tgl_nbds[-1]} · O2O {R['tglO2O']}")
    print("Selesai. Naikkan angka ?v= pada file data di index.html agar browser memuat data baru.")


if __name__ == "__main__":
    main()
